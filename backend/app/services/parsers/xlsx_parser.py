from io import BytesIO

import openpyxl

from app.services.parsers import ParsedQuestion
from app.services.parsers._type_detection import detect_question_type

_QUESTION_HEADER_KEYWORDS = ("question",)
_ANSWER_HEADER_KEYWORDS = ("answer", "response")


def _find_header_columns(row_cells) -> tuple[int | None, int | None]:
    question_col: int | None = None
    answer_col: int | None = None
    for cell in row_cells:
        if not isinstance(cell.value, str):
            continue
        value = cell.value.strip().lower()
        if question_col is None and any(kw in value for kw in _QUESTION_HEADER_KEYWORDS):
            question_col = cell.column
        if answer_col is None and any(kw in value for kw in _ANSWER_HEADER_KEYWORDS):
            answer_col = cell.column
    return question_col, answer_col


def parse(content: bytes) -> list[ParsedQuestion]:
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    results: list[ParsedQuestion] = []

    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue

        header_row = next(ws.iter_rows(min_row=1, max_row=1), [])
        question_col, answer_col = _find_header_columns(header_row)
        has_header = question_col is not None

        if question_col is None:
            question_col = 1  # column A
        if answer_col is None:
            answer_col = question_col + 1

        start_row = 2 if has_header else 1

        for row in ws.iter_rows(min_row=start_row):
            cell = row[question_col - 1] if len(row) >= question_col else None
            if cell is None or not isinstance(cell.value, str) or not cell.value.strip():
                continue
            text = cell.value.strip()

            results.append(
                ParsedQuestion(
                    text=text,
                    position={
                        "kind": "xlsx",
                        "sheet_name": ws.title,
                        "question_row": cell.row,
                        "question_col": question_col,
                        "answer_row": cell.row,
                        "answer_col": answer_col,
                    },
                    detected_type=detect_question_type(text),
                    section=ws.title,
                )
            )

    return results

from io import BytesIO

from docx import Document

from app.services.parsers import ParsedQuestion
from app.services.parsers._type_detection import detect_question_type

_PLACEHOLDER_MAX_LEN = 3


def _fingerprint(text: str) -> str:
    return text.strip()[:80]


def parse(content: bytes) -> list[ParsedQuestion]:
    doc = Document(BytesIO(content))
    results: list[ParsedQuestion] = []

    paragraphs = doc.paragraphs
    for i, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text.endswith("?"):
            continue

        next_text = paragraphs[i + 1].text.strip() if i + 1 < len(paragraphs) else None
        has_placeholder = next_text is not None and len(next_text) <= _PLACEHOLDER_MAX_LEN

        results.append(
            ParsedQuestion(
                text=text,
                position={
                    "kind": "docx",
                    "anchor": "paragraph",
                    "paragraph_index": (i + 1) if has_placeholder else i,
                    "has_placeholder": has_placeholder,
                    "question_fingerprint": _fingerprint(text),
                },
                detected_type=detect_question_type(text),
                section=None,
            )
        )

    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            if len(row.cells) < 2:
                continue
            question_text = row.cells[0].text.strip()
            if not question_text.endswith("?"):
                continue

            results.append(
                ParsedQuestion(
                    text=question_text,
                    position={
                        "kind": "docx",
                        "anchor": "table_cell",
                        "table_index": table_index,
                        "row_index": row_index,
                        "col_index": 1,
                        "has_placeholder": True,
                        "question_fingerprint": _fingerprint(question_text),
                    },
                    detected_type=detect_question_type(question_text),
                    section=None,
                )
            )

    return results

from io import BytesIO

from docx import Document

from app.services.kb_chunking import KbChunk, _split_by_max_chars


def _is_heading(paragraph) -> bool:
    style_name = (paragraph.style.name or "").lower()
    if style_name.startswith("heading") or style_name == "title":
        return True
    text = paragraph.text.strip()
    if not text or len(text) > 100:
        return False
    bold_runs = [r for r in paragraph.runs if r.text.strip()]
    return bool(bold_runs) and all(r.bold for r in bold_runs)


def _flush_section(chunks: list[KbChunk], section: str | None, body_lines: list[str]) -> None:
    body = "\n".join(body_lines).strip()
    if not body:
        return
    for piece in _split_by_max_chars(body):
        chunks.append(KbChunk(text=piece, page=None, section=section))


def chunk(content: bytes) -> list[KbChunk]:
    doc = Document(BytesIO(content))
    chunks: list[KbChunk] = []

    current_section: str | None = None
    current_lines: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if _is_heading(paragraph):
            _flush_section(chunks, current_section, current_lines)
            current_section, current_lines = text, []
        else:
            current_lines.append(text)
    _flush_section(chunks, current_section, current_lines)

    for i, table in enumerate(doc.tables):
        rows_text = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        body = "\n".join(r for r in rows_text if r.strip())
        if not body:
            continue
        for piece in _split_by_max_chars(body):
            chunks.append(KbChunk(text=piece, page=None, section=f"Table {i + 1}"))

    return chunks

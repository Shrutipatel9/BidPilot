import asyncio
import copy
import csv
import io
from io import BytesIO

import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.models.answer import Answer
from app.models.question import Question


def _resolve_writable_cell(ws, row: int, col: int):
    cell = ws.cell(row=row, column=col)
    if type(cell).__name__ != "MergedCell":
        return cell
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col)
    return cell


def _fill_xlsx_sync(original_bytes: bytes, fills: list[tuple[dict, str]]) -> bytes:
    wb = openpyxl.load_workbook(BytesIO(original_bytes), data_only=False)
    for position, answer_text in fills:
        ws = wb[position["sheet_name"]]
        cell = _resolve_writable_cell(ws, position["answer_row"], position["answer_col"])
        cell.value = answer_text
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def fill_xlsx(original_bytes: bytes, fills: list[tuple[dict, str]]) -> bytes:
    return await asyncio.to_thread(_fill_xlsx_sync, original_bytes, fills)


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for extra in runs[1:]:
        extra.text = ""


def _set_cell_text(cell, text: str) -> None:
    _set_paragraph_text(cell.paragraphs[0], text)
    for extra_paragraph in cell.paragraphs[1:]:
        extra_paragraph._element.getparent().remove(extra_paragraph._element)


def _insert_answer_paragraph_after(doc: Document, question_paragraph: Paragraph, text: str) -> None:
    new_element = copy.deepcopy(question_paragraph._element)
    for run_element in new_element.findall(qn("w:r")):
        new_element.remove(run_element)
    question_paragraph._element.addnext(new_element)
    new_paragraph = Paragraph(new_element, question_paragraph._parent)
    new_paragraph.style = doc.styles["Normal"]
    new_paragraph.add_run(text)


def _fingerprint(text: str) -> str:
    return text.strip()[:80]


def _fill_docx_sync(original_bytes: bytes, fills: list[tuple[dict, str]]) -> tuple[bytes, list[str]]:
    doc = Document(BytesIO(original_bytes))
    skipped_fingerprints: list[str] = []

    for position, answer_text in fills:
        if position["anchor"] == "table_cell":
            table = doc.tables[position["table_index"]]
            cell = table.rows[position["row_index"]].cells[position["col_index"]]
            question_text = table.rows[position["row_index"]].cells[0].text
            if _fingerprint(question_text) != position["question_fingerprint"]:
                skipped_fingerprints.append(position["question_fingerprint"])
                continue
            _set_cell_text(cell, answer_text)
            continue

        paragraphs = doc.paragraphs
        target_idx = position["paragraph_index"]
        question_idx = target_idx - 1 if position["has_placeholder"] else target_idx
        if (
            question_idx < 0
            or question_idx >= len(paragraphs)
            or _fingerprint(paragraphs[question_idx].text) != position["question_fingerprint"]
        ):
            skipped_fingerprints.append(position["question_fingerprint"])
            continue

        if position["has_placeholder"]:
            _set_paragraph_text(paragraphs[target_idx], answer_text)
        else:
            _insert_answer_paragraph_after(doc, paragraphs[question_idx], answer_text)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue(), skipped_fingerprints


async def fill_docx(original_bytes: bytes, fills: list[tuple[dict, str]]) -> tuple[bytes, list[str]]:
    return await asyncio.to_thread(_fill_docx_sync, original_bytes, fills)


def export_csv(questions: list[Question], answers_by_question_id: dict) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Question", "Section", "Answer", "Status", "Confidence"])
    for question in questions:
        answer: Answer | None = answers_by_question_id.get(question.id)
        writer.writerow(
            [
                question.text,
                question.section or "",
                answer.text if answer else "",
                answer.status.value if answer else "not_started",
                answer.confidence if answer and answer.confidence is not None else "",
            ]
        )
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 CSVs without mangling

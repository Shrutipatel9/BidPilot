import io
from unittest.mock import AsyncMock, patch

import openpyxl

from app.core.llm_provider import DraftAnswer, LLMDraftError


async def _signup(client, email, password="correct-horse-battery"):
    resp = await client.post("/api/auth/signup", json={"email": email, "password": password})
    return resp.json()["access_token"]


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


async def _create_org(client, token, name="Acme Inc"):
    resp = await client.post("/api/orgs", json={"name": name}, headers=_auth(token))
    return resp.json()["id"]


def _build_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Question"
    ws["B1"] = "Answer"
    ws["A2"] = "Do you encrypt data at rest?"
    ws["A3"] = "Do you maintain a SOC 2 report (Yes/No)?"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _setup_confirmed_project(client, org_id, token):
    resp = await client.post(
        f"/api/orgs/{org_id}/projects",
        files={"file": ("q.xlsx", io.BytesIO(_build_xlsx_bytes()), "application/octet-stream")},
        data={"name": "Test Project"},
        headers=_auth(token),
    )
    project_id = resp.json()["id"]
    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/parse", headers=_auth(token))
    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/questions/confirm", headers=_auth(token))
    return project_id


@patch("app.services.drafting_service.generate_draft_answer", new_callable=AsyncMock)
async def test_drafting_creates_answers_and_revisions(mock_draft, client):
    mock_draft.return_value = DraftAnswer(
        answer_text="Yes, we encrypt all data at rest using AES-256.",
        confidence=90,
        choice="Yes",
        needs_review=False,
    )

    token = await _signup(client, "owner1@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_confirmed_project(client, org_id, token)

    resp = await client.post(f"/api/orgs/{org_id}/projects/{project_id}/draft", headers=_auth(token))
    assert resp.status_code == 202

    resp = await client.get(f"/api/orgs/{org_id}/projects/{project_id}/answers", headers=_auth(token))
    assert resp.status_code == 200
    answers = resp.json()
    assert len(answers) == 2
    for answer in answers:
        assert answer["status"] == "ai_drafted"
        assert answer["confidence"] == 90
        assert answer["text"] == "Yes, we encrypt all data at rest using AES-256."

    resp = await client.get(f"/api/orgs/{org_id}/projects/{project_id}", headers=_auth(token))
    assert resp.json()["status"] == "drafted"


@patch("app.services.drafting_service.generate_draft_answer", new_callable=AsyncMock)
async def test_drafting_flags_needs_review_when_llm_says_so(mock_draft, client):
    mock_draft.return_value = DraftAnswer(
        answer_text="We believe so, but cannot confirm without more information.",
        confidence=30,
        choice=None,
        needs_review=True,
    )

    token = await _signup(client, "owner2@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_confirmed_project(client, org_id, token)

    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/draft", headers=_auth(token))

    resp = await client.get(f"/api/orgs/{org_id}/projects/{project_id}/answers", headers=_auth(token))
    assert all(a["status"] == "needs_review" for a in resp.json())


@patch("app.services.drafting_service.generate_draft_answer", new_callable=AsyncMock)
async def test_drafting_failure_does_not_abort_whole_batch(mock_draft, client):
    # First question fails outright, second succeeds — the batch must still complete both.
    mock_draft.side_effect = [LLMDraftError("provider down"), DraftAnswer(
        answer_text="Yes.", confidence=70, choice="Yes", needs_review=False
    )]

    token = await _signup(client, "owner3@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_confirmed_project(client, org_id, token)

    resp = await client.post(f"/api/orgs/{org_id}/projects/{project_id}/draft", headers=_auth(token))
    assert resp.status_code == 202

    resp = await client.get(f"/api/orgs/{org_id}/projects/{project_id}/answers", headers=_auth(token))
    statuses = {a["status"] for a in resp.json()}
    assert statuses == {"needs_review", "ai_drafted"}

    resp = await client.get(f"/api/orgs/{org_id}/projects/{project_id}", headers=_auth(token))
    assert resp.json()["status"] == "drafted"  # batch still completes despite one failure


async def test_drafting_tenant_isolation(client):
    token_a = await _signup(client, "orga@example.com")
    org_a = await _create_org(client, token_a, "Org A")
    project_id = await _setup_confirmed_project(client, org_a, token_a)

    token_b = await _signup(client, "orgb@example.com")
    org_b = await _create_org(client, token_b, "Org B")

    resp = await client.post(f"/api/orgs/{org_a}/projects/{project_id}/draft", headers=_auth(token_b))
    assert resp.status_code == 403

    resp = await client.get(f"/api/orgs/{org_a}/projects/{project_id}/answers", headers=_auth(token_b))
    assert resp.status_code == 403

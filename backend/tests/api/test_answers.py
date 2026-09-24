import io
import uuid

import openpyxl
from sqlalchemy import select

from app.models.answer import Answer, AnswerStatus
from app.models.membership import Membership, MembershipRole
from app.models.question import Question


async def _signup(client, email, password="correct-horse-battery"):
    resp = await client.post("/api/auth/signup", json={"email": email, "password": password})
    return resp.json()["access_token"]


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


async def _create_org(client, token, name="Acme Inc"):
    resp = await client.post("/api/orgs", json={"name": name}, headers=_auth(token))
    return resp.json()["id"]


def _build_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Question"
    ws["B1"] = "Answer"
    ws["A2"] = "Do you encrypt data at rest?"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _setup_drafted_answer(client, db_session, org_id, token) -> tuple[str, str]:
    """Creates a confirmed project with one question, then fast-forwards its stub answer
    straight to ai_drafted (bypassing a real/mocked LLM call — 1.4 tests the review workflow,
    not drafting itself, which is already covered by test_drafting.py)."""
    resp = await client.post(
        f"/api/orgs/{org_id}/projects",
        files={"file": ("q.xlsx", io.BytesIO(_build_xlsx_bytes()), "application/octet-stream")},
        data={"name": "Test Project"},
        headers=_auth(token),
    )
    project_id = resp.json()["id"]
    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/parse", headers=_auth(token))
    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/questions/confirm", headers=_auth(token))

    answer = await db_session.scalar(
        select(Answer)
        .join(Question, Question.id == Answer.question_id)
        .where(Question.project_id == uuid.UUID(project_id))
    )
    answer.status = AnswerStatus.ai_drafted
    answer.text = "Yes, we encrypt all data at rest."
    answer.confidence = 90
    await db_session.commit()
    await db_session.refresh(answer)
    return project_id, str(answer.id)


async def _add_member(client, db_session, org_id, email, role):
    token = await _signup(client, email)
    me = await client.get("/api/auth/me", headers=_auth(token))
    db_session.add(Membership(user_id=uuid.UUID(me.json()["id"]), org_id=uuid.UUID(org_id), role=role))
    await db_session.commit()
    return token


async def test_edit_answer_creates_revision_and_bumps_version(client, db_session):
    token = await _signup(client, "owner1@example.com")
    org_id = await _create_org(client, token)
    project_id, answer_id = await _setup_drafted_answer(client, db_session, org_id, token)

    resp = await client.patch(
        f"/api/orgs/{org_id}/projects/{project_id}/answers/{answer_id}",
        json={"text": "Yes, AES-256 encryption at rest, verified annually."},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["text"] == "Yes, AES-256 encryption at rest, verified annually."
    assert body["version"] == 2
    assert body["status"] == "in_review"


async def test_approve_and_reject_transitions(client, db_session):
    token = await _signup(client, "owner2@example.com")
    org_id = await _create_org(client, token)
    project_id, answer_id = await _setup_drafted_answer(client, db_session, org_id, token)

    resp = await client.post(
        f"/api/orgs/{org_id}/projects/{project_id}/answers/{answer_id}/approve", headers=_auth(token)
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["status"] == "approved"

    project_id2, answer_id2 = await _setup_drafted_answer(client, db_session, org_id, token)
    resp = await client.post(
        f"/api/orgs/{org_id}/projects/{project_id2}/answers/{answer_id2}/reject",
        json={"reason": "Needs a more specific answer."},
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["status"] == "rejected"


async def test_viewer_cannot_edit_or_approve(client, db_session):
    owner_token = await _signup(client, "owner3@example.com")
    org_id = await _create_org(client, owner_token)
    project_id, answer_id = await _setup_drafted_answer(client, db_session, org_id, owner_token)

    viewer_token = await _add_member(client, db_session, org_id, "viewer3@example.com", MembershipRole.viewer)

    resp = await client.patch(
        f"/api/orgs/{org_id}/projects/{project_id}/answers/{answer_id}",
        json={"text": "Should fail"},
        headers=_auth(viewer_token),
    )
    assert resp.status_code == 403

    resp = await client.post(
        f"/api/orgs/{org_id}/projects/{project_id}/answers/{answer_id}/approve", headers=_auth(viewer_token)
    )
    assert resp.status_code == 403


async def test_knowledge_manager_cannot_approve(client, db_session):
    owner_token = await _signup(client, "owner4@example.com")
    org_id = await _create_org(client, owner_token)
    project_id, answer_id = await _setup_drafted_answer(client, db_session, org_id, owner_token)

    km_token = await _add_member(
        client, db_session, org_id, "km4@example.com", MembershipRole.knowledge_manager
    )

    resp = await client.post(
        f"/api/orgs/{org_id}/projects/{project_id}/answers/{answer_id}/approve", headers=_auth(km_token)
    )
    assert resp.status_code == 403


async def test_reviewer_can_approve_but_not_edit(client, db_session):
    owner_token = await _signup(client, "owner5@example.com")
    org_id = await _create_org(client, owner_token)
    project_id, answer_id = await _setup_drafted_answer(client, db_session, org_id, owner_token)

    reviewer_token = await _add_member(
        client, db_session, org_id, "reviewer5@example.com", MembershipRole.reviewer
    )

    resp = await client.patch(
        f"/api/orgs/{org_id}/projects/{project_id}/answers/{answer_id}",
        json={"text": "Should fail — reviewer can't edit"},
        headers=_auth(reviewer_token),
    )
    assert resp.status_code == 403

    resp = await client.post(
        f"/api/orgs/{org_id}/projects/{project_id}/answers/{answer_id}/approve", headers=_auth(reviewer_token)
    )
    assert resp.status_code == 200, resp.text


async def test_answer_tenant_isolation(client, db_session):
    token_a = await _signup(client, "orga@example.com")
    org_a = await _create_org(client, token_a, "Org A")
    project_id, answer_id = await _setup_drafted_answer(client, db_session, org_a, token_a)

    token_b = await _signup(client, "orgb@example.com")
    org_b = await _create_org(client, token_b, "Org B")

    resp = await client.patch(
        f"/api/orgs/{org_a}/projects/{project_id}/answers/{answer_id}",
        json={"text": "cross-tenant edit attempt"},
        headers=_auth(token_b),
    )
    assert resp.status_code == 403

    resp = await client.post(
        f"/api/orgs/{org_a}/projects/{project_id}/answers/{answer_id}/approve", headers=_auth(token_b)
    )
    assert resp.status_code == 403

    resp = await client.post(
        f"/api/orgs/{org_a}/projects/{project_id}/answers/{answer_id}/reject",
        json={"reason": "cross-tenant"},
        headers=_auth(token_b),
    )
    assert resp.status_code == 403


async def test_answer_scoping_rejects_foreign_project_even_for_a_real_member(client, db_session):
    """The 403 tests above only prove the membership gate works (token_b has no membership in
    org_a at all). This proves the actual dual-filter/transitive-scoping logic: a genuine member
    of org_b, authenticated with their OWN org_id in the path, must still be refused when the
    project_id/answer_id in the URL belong to org_a — this has to 404 (get_project not finding a
    match), not merely rely on a 403 that never reaches that query."""
    token_a = await _signup(client, "orga2@example.com")
    org_a = await _create_org(client, token_a, "Org A")
    project_id, answer_id = await _setup_drafted_answer(client, db_session, org_a, token_a)

    token_b = await _signup(client, "orgb2@example.com")
    org_b = await _create_org(client, token_b, "Org B")

    resp = await client.patch(
        f"/api/orgs/{org_b}/projects/{project_id}/answers/{answer_id}",
        json={"text": "cross-tenant edit via own org_id"},
        headers=_auth(token_b),
    )
    assert resp.status_code == 404

    resp = await client.post(
        f"/api/orgs/{org_b}/projects/{project_id}/answers/{answer_id}/approve", headers=_auth(token_b)
    )
    assert resp.status_code == 404

    resp = await client.post(
        f"/api/orgs/{org_b}/projects/{project_id}/answers/{answer_id}/reject",
        json={"reason": "cross-tenant via own org_id"},
        headers=_auth(token_b),
    )
    assert resp.status_code == 404

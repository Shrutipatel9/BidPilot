import io
import uuid
from unittest.mock import AsyncMock, patch

import openpyxl

from app.core.llm_provider import DraftAnswer
from app.models.knowledge_chunk import KnowledgeChunk
from app.models.knowledge_document import DocumentStatus, KnowledgeDocument
from app.services.retrieval_service import RetrievedChunk


def _vector(dim_value_pairs, dim=768):
    v = [0.0] * dim
    for i, val in dim_value_pairs:
        v[i] = val
    return v


async def _signup(client, email, password="correct-horse-battery"):
    resp = await client.post("/api/auth/signup", json={"email": email, "password": password})
    return resp.json()["access_token"]


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


async def _create_org(client, token, name="Acme Inc"):
    resp = await client.post("/api/orgs", json={"name": name}, headers=_auth(token))
    return resp.json()["id"]


def _build_xlsx_bytes(question="Do you encrypt data at rest?") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Question"
    ws["B1"] = "Answer"
    ws["A2"] = question
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _setup_confirmed_project(client, org_id, token, question="Do you encrypt data at rest?"):
    resp = await client.post(
        f"/api/orgs/{org_id}/projects",
        files={"file": ("q.xlsx", io.BytesIO(_build_xlsx_bytes(question)), "application/octet-stream")},
        data={"name": "Test Project"},
        headers=_auth(token),
    )
    project_id = resp.json()["id"]
    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/parse", headers=_auth(token))
    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/questions/confirm", headers=_auth(token))
    return project_id


@patch("app.services.drafting_service.retrieve_top_chunks", new_callable=AsyncMock)
async def test_zero_chunks_short_circuits_without_calling_the_llm(mock_retrieve, client):
    mock_retrieve.return_value = []

    with patch("app.services.drafting_service.generate_draft_answer", new_callable=AsyncMock) as mock_draft:
        token = await _signup(client, "owner1@example.com")
        org_id = await _create_org(client, token)
        project_id = await _setup_confirmed_project(client, org_id, token)

        await client.post(f"/api/orgs/{org_id}/projects/{project_id}/draft", headers=_auth(token))

        mock_draft.assert_not_called()

    resp = await client.get(f"/api/orgs/{org_id}/projects/{project_id}/answers", headers=_auth(token))
    answers = resp.json()
    assert len(answers) == 1
    assert answers[0]["text"] == "Insufficient information"
    assert answers[0]["confidence"] == 0
    assert answers[0]["status"] == "needs_review"


@patch("app.services.drafting_service.retrieve_top_chunks", new_callable=AsyncMock)
@patch("app.services.drafting_service.generate_draft_answer", new_callable=AsyncMock)
async def test_grounded_answer_gets_citations_and_blended_confidence(mock_draft, mock_retrieve, client):
    chunk = RetrievedChunk(
        chunk_id=uuid.uuid4(),
        document_id=uuid.uuid4(),
        document_title="Security Policy",
        page=3,
        section=None,
        text="We encrypt all customer data at rest using AES-256 encryption keys rotated every 90 days.",
        similarity=0.55,  # deliberately lower than the LLM's own self-assessed confidence
    )
    mock_retrieve.return_value = [chunk]
    mock_draft.return_value = DraftAnswer(
        answer_text="Yes, AES-256 at rest.", confidence=95, choice="Yes", needs_review=False
    )

    token = await _signup(client, "owner2@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_confirmed_project(client, org_id, token)

    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/draft", headers=_auth(token))

    resp = await client.get(f"/api/orgs/{org_id}/projects/{project_id}/answers", headers=_auth(token))
    answer = resp.json()[0]
    assert answer["status"] == "ai_drafted"
    # Confidence is capped by retrieval quality (55%), not the LLM's own 95% self-assessment.
    assert answer["confidence"] == 55
    assert len(answer["citations"]) == 1
    citation = answer["citations"][0]
    assert citation["document_title"] == "Security Policy"
    assert citation["page"] == 3
    assert citation["snippet"].startswith("We encrypt all customer data at rest")


@patch("app.services.drafting_service.retrieve_top_chunks", new_callable=AsyncMock)
@patch("app.services.drafting_service.generate_draft_answer", new_callable=AsyncMock)
async def test_llm_judged_insufficiency_clears_citations_despite_retrieved_chunks(mock_draft, mock_retrieve, client):
    chunk = RetrievedChunk(
        chunk_id=uuid.uuid4(),
        document_id=uuid.uuid4(),
        document_title="Unrelated Policy",
        page=None,
        section="HR",
        text="Employees receive 20 days of paid time off per year.",
        similarity=0.4,
    )
    mock_retrieve.return_value = [chunk]
    mock_draft.return_value = DraftAnswer(
        answer_text="Insufficient information", confidence=20, choice=None, needs_review=True
    )

    token = await _signup(client, "owner3@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_confirmed_project(client, org_id, token)

    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/draft", headers=_auth(token))

    resp = await client.get(f"/api/orgs/{org_id}/projects/{project_id}/answers", headers=_auth(token))
    answer = resp.json()[0]
    assert answer["text"] == "Insufficient information"
    assert answer["status"] == "needs_review"
    assert answer["citations"] == []


@patch("app.services.retrieval_service.generate_embedding", new_callable=AsyncMock)
@patch("app.services.drafting_service.generate_draft_answer", new_callable=AsyncMock)
async def test_drafting_retrieval_never_crosses_tenant_boundary_end_to_end(mock_draft, mock_embed, client, db_session):
    """Unlike the other tests in this file (which mock retrieve_top_chunks directly), this
    exercises the REAL pgvector query through the full route -> run_drafting_for_project ->
    retrieve_top_chunks path, with real cross-org KnowledgeChunk rows seeded in the same
    database — the scenario the retrieval_service unit tests alone can't prove end to end."""
    mock_embed.return_value = _vector([(0, 1.0)])
    mock_draft.return_value = DraftAnswer(answer_text="Yes.", confidence=80, choice="Yes", needs_review=False)

    token_a = await _signup(client, "orga_e2e@example.com")
    org_a = await _create_org(client, token_a, "Org A E2E")
    project_id = await _setup_confirmed_project(client, org_a, token_a, question="Do you encrypt data?")

    token_b = await _signup(client, "orgb_e2e@example.com")
    org_b = await _create_org(client, token_b, "Org B E2E")

    doc_a = KnowledgeDocument(
        org_id=uuid.UUID(org_a), title="Org A Doc", source_file_key="a", source_file_ext="txt",
        status=DocumentStatus.indexed,
    )
    doc_b = KnowledgeDocument(
        org_id=uuid.UUID(org_b), title="Org B Doc", source_file_key="b", source_file_ext="txt",
        status=DocumentStatus.indexed,
    )
    db_session.add_all([doc_a, doc_b])
    await db_session.flush()

    # Org B's chunk is a near-perfect match for the query embedding — it would outrank org A's
    # own weak match if the org_id filter were missing or applied after the fact.
    db_session.add_all(
        [
            KnowledgeChunk(
                document_id=doc_a.id, org_id=uuid.UUID(org_a), text="org a weak content",
                embedding=_vector([(5, 1.0)]),
            ),
            KnowledgeChunk(
                document_id=doc_b.id, org_id=uuid.UUID(org_b), text="org b perfect content",
                embedding=_vector([(0, 1.0)]),
            ),
        ]
    )
    await db_session.commit()

    await client.post(f"/api/orgs/{org_a}/projects/{project_id}/draft", headers=_auth(token_a))

    resp = await client.get(f"/api/orgs/{org_a}/projects/{project_id}/answers", headers=_auth(token_a))
    answer = resp.json()[0]
    assert len(answer["citations"]) == 1
    assert answer["citations"][0]["document_id"] == str(doc_a.id)
    assert answer["citations"][0]["document_title"] == "Org A Doc"

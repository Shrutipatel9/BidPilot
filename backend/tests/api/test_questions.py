import io

import openpyxl
from docx import Document


async def _signup(client, email, password="correct-horse-battery"):
    resp = await client.post("/api/auth/signup", json={"email": email, "password": password})
    return resp.json()["access_token"]


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


async def _create_org(client, token, name="Acme Inc"):
    resp = await client.post("/api/orgs", json={"name": name}, headers=_auth(token))
    return resp.json()["id"]


def _build_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Question"
    ws["B1"] = "Answer"
    ws["A2"] = "Do you encrypt data at rest?"
    ws["A3"] = "How many employees does your company have?"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Section 1: Security")
    doc.add_paragraph("Do you have a SOC 2 report?")
    doc.add_paragraph("")  # placeholder answer line
    doc.add_paragraph("Do you encrypt data in transit (Yes/No)?")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _create_project_with_file(client, org_id, token, filename, content, content_type):
    resp = await client.post(
        f"/api/orgs/{org_id}/projects",
        files={"file": (filename, io.BytesIO(content), content_type)},
        data={"name": "Test Project"},
        headers=_auth(token),
    )
    return resp.json()["id"]


async def test_parse_xlsx_project(client):
    token = await _signup(client, "owner1@example.com")
    org_id = await _create_org(client, token)
    project_id = await _create_project_with_file(
        client, org_id, token, "q.xlsx", _build_xlsx_bytes(), "application/octet-stream"
    )

    resp = await client.post(f"/api/orgs/{org_id}/projects/{project_id}/parse", headers=_auth(token))
    assert resp.status_code == 200, resp.text
    questions = resp.json()
    assert len(questions) == 2
    texts = {q["text"] for q in questions}
    assert "Do you encrypt data at rest?" in texts
    numeric_q = next(q for q in questions if "employees" in q["text"])
    assert numeric_q["type"] == "numeric"
    assert questions[0]["position"]["kind"] == "xlsx"


async def test_parse_docx_project(client):
    token = await _signup(client, "owner2@example.com")
    org_id = await _create_org(client, token)
    project_id = await _create_project_with_file(
        client, org_id, token, "q.docx", _build_docx_bytes(), "application/octet-stream"
    )

    resp = await client.post(f"/api/orgs/{org_id}/projects/{project_id}/parse", headers=_auth(token))
    assert resp.status_code == 200, resp.text
    questions = resp.json()
    assert len(questions) == 2
    yes_no = next(q for q in questions if "transit" in q["text"])
    assert yes_no["type"] == "yes_no"
    placeholder_q = next(q for q in questions if "SOC 2" in q["text"])
    assert placeholder_q["position"]["has_placeholder"] is True


async def test_update_and_confirm_questions(client):
    token = await _signup(client, "owner3@example.com")
    org_id = await _create_org(client, token)
    project_id = await _create_project_with_file(
        client, org_id, token, "q.xlsx", _build_xlsx_bytes(), "application/octet-stream"
    )
    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/parse", headers=_auth(token))

    resp = await client.get(f"/api/orgs/{org_id}/projects/{project_id}/questions", headers=_auth(token))
    question_id = resp.json()[0]["id"]

    resp = await client.patch(
        f"/api/orgs/{org_id}/projects/{project_id}/questions/{question_id}",
        json={"text": "Corrected question text?"},
        headers=_auth(token),
    )
    assert resp.status_code == 200
    assert resp.json()["text"] == "Corrected question text?"

    resp = await client.post(f"/api/orgs/{org_id}/projects/{project_id}/questions/confirm", headers=_auth(token))
    assert resp.status_code == 200
    assert resp.json()["status"] == "questions_confirmed"

    # Confirming again should be idempotent (no duplicate Answer rows), not error.
    resp = await client.post(f"/api/orgs/{org_id}/projects/{project_id}/questions/confirm", headers=_auth(token))
    assert resp.status_code == 200


async def test_confirm_without_parsing_rejected(client):
    token = await _signup(client, "owner4@example.com")
    org_id = await _create_org(client, token)
    project_id = await _create_project_with_file(
        client, org_id, token, "q.xlsx", _build_xlsx_bytes(), "application/octet-stream"
    )

    resp = await client.post(f"/api/orgs/{org_id}/projects/{project_id}/questions/confirm", headers=_auth(token))
    assert resp.status_code == 400


async def test_questions_tenant_isolation(client):
    token_a = await _signup(client, "orga@example.com")
    org_a = await _create_org(client, token_a, "Org A")
    project_id = await _create_project_with_file(
        client, org_a, token_a, "q.xlsx", _build_xlsx_bytes(), "application/octet-stream"
    )
    await client.post(f"/api/orgs/{org_a}/projects/{project_id}/parse", headers=_auth(token_a))

    token_b = await _signup(client, "orgb@example.com")
    org_b = await _create_org(client, token_b, "Org B")

    resp = await client.get(f"/api/orgs/{org_a}/projects/{project_id}/questions", headers=_auth(token_b))
    assert resp.status_code == 403

    resp = await client.post(f"/api/orgs/{org_a}/projects/{project_id}/parse", headers=_auth(token_b))
    assert resp.status_code == 403

    questions = (
        await client.get(f"/api/orgs/{org_a}/projects/{project_id}/questions", headers=_auth(token_a))
    ).json()
    question_id = questions[0]["id"]

    resp = await client.patch(
        f"/api/orgs/{org_a}/projects/{project_id}/questions/{question_id}",
        json={"text": "hijacked"},
        headers=_auth(token_b),
    )
    assert resp.status_code == 403

    resp = await client.post(f"/api/orgs/{org_a}/projects/{project_id}/questions/confirm", headers=_auth(token_b))
    assert resp.status_code == 403

import csv
import io
import uuid

import openpyxl
from docx import Document
from sqlalchemy import select

from app.models.answer import Answer, AnswerStatus
from app.models.question import Question


async def _signup(client, email, password="correct-horse-battery"):
    resp = await client.post("/api/auth/signup", json={"email": email, "password": password})
    return resp.json()["access_token"]


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


async def _create_org(client, token, name="Acme Inc"):
    resp = await client.post("/api/orgs", json={"name": name}, headers=_auth(token))
    return resp.json()["id"]


def _build_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Question"
    ws["B1"] = "Answer"
    ws["A2"] = "Do you encrypt data at rest?"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Do you have a SOC 2 report?")
    doc.add_paragraph("")  # placeholder answer line
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _setup_project_with_answer(client, db_session, org_id, token, filename, content, answer_text):
    resp = await client.post(
        f"/api/orgs/{org_id}/projects",
        files={"file": (filename, io.BytesIO(content), "application/octet-stream")},
        data={"name": "Export Test Project"},
        headers=_auth(token),
    )
    project_id = resp.json()["id"]
    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/parse", headers=_auth(token))
    await client.post(f"/api/orgs/{org_id}/projects/{project_id}/questions/confirm", headers=_auth(token))

    answer = await db_session.scalar(
        select(Answer)
        .join(Question, Question.id == Answer.question_id)
        .where(Question.project_id == uuid.UUID(project_id))
    )
    answer.text = answer_text
    answer.status = AnswerStatus.ai_drafted
    await db_session.commit()
    return project_id


async def test_export_xlsx_fills_correct_cell(client, db_session):
    token = await _signup(client, "owner1@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_project_with_answer(
        client, db_session, org_id, token, "q.xlsx", _build_xlsx_bytes(), "Yes, AES-256 at rest."
    )

    resp = await client.get(
        f"/api/orgs/{org_id}/projects/{project_id}/export?format=xlsx", headers=_auth(token)
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws["A2"].value == "Do you encrypt data at rest?"  # original question cell untouched
    assert ws["B2"].value == "Yes, AES-256 at rest."


async def test_export_docx_fills_placeholder_paragraph(client, db_session):
    token = await _signup(client, "owner2@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_project_with_answer(
        client, db_session, org_id, token, "q.docx", _build_docx_bytes(), "Yes, SOC 2 Type II attached."
    )

    resp = await client.get(
        f"/api/orgs/{org_id}/projects/{project_id}/export?format=docx", headers=_auth(token)
    )
    assert resp.status_code == 200, resp.text

    doc = Document(io.BytesIO(resp.content))
    texts = [p.text for p in doc.paragraphs]
    assert "Do you have a SOC 2 report?" in texts
    assert "Yes, SOC 2 Type II attached." in texts


async def test_export_docx_inserts_paragraph_when_no_placeholder(client, db_session):
    # No blank line after the question — the parser records has_placeholder=False, so export
    # must clone/insert a fresh paragraph after the question rather than overwrite anything.
    doc = Document()
    doc.add_paragraph("Do you encrypt data in transit?")
    doc.add_paragraph("Next section starts here.")
    buf = io.BytesIO()
    doc.save(buf)

    token = await _signup(client, "owner5@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_project_with_answer(
        client, db_session, org_id, token, "q.docx", buf.getvalue(), "Yes, TLS 1.2+ in transit."
    )

    resp = await client.get(
        f"/api/orgs/{org_id}/projects/{project_id}/export?format=docx", headers=_auth(token)
    )
    assert resp.status_code == 200, resp.text

    result_doc = Document(io.BytesIO(resp.content))
    texts = [p.text for p in result_doc.paragraphs]
    assert "Do you encrypt data in transit?" in texts
    assert "Yes, TLS 1.2+ in transit." in texts
    assert "Next section starts here." in texts
    question_idx = texts.index("Do you encrypt data in transit?")
    assert texts[question_idx + 1] == "Yes, TLS 1.2+ in transit."


async def test_export_xlsx_writes_through_merged_answer_cell(client, db_session):
    # If the answer cell is part of a merged range, only the anchor (top-left) cell is writable —
    # openpyxl raises on writing to a MergedCell directly, so export must resolve to the anchor.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Question"
    ws["B1"] = "Answer"
    ws["A2"] = "Do you encrypt data at rest?"
    ws.merge_cells("B2:C2")
    buf = io.BytesIO()
    wb.save(buf)

    token = await _signup(client, "owner6@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_project_with_answer(
        client, db_session, org_id, token, "q.xlsx", buf.getvalue(), "Yes, merged cell answer."
    )

    resp = await client.get(
        f"/api/orgs/{org_id}/projects/{project_id}/export?format=xlsx", headers=_auth(token)
    )
    assert resp.status_code == 200, resp.text

    result_wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    result_ws = result_wb.active
    assert result_ws["B2"].value == "Yes, merged cell answer."


async def test_export_csv_structure(client, db_session):
    token = await _signup(client, "owner3@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_project_with_answer(
        client, db_session, org_id, token, "q.xlsx", _build_xlsx_bytes(), "Yes."
    )

    resp = await client.get(
        f"/api/orgs/{org_id}/projects/{project_id}/export?format=csv", headers=_auth(token)
    )
    assert resp.status_code == 200, resp.text
    rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))
    assert rows[0] == ["Question", "Section", "Answer", "Status", "Confidence"]
    assert rows[1][0] == "Do you encrypt data at rest?"
    assert rows[1][2] == "Yes."
    assert rows[1][3] == "ai_drafted"


async def test_export_format_must_match_source(client, db_session):
    token = await _signup(client, "owner4@example.com")
    org_id = await _create_org(client, token)
    project_id = await _setup_project_with_answer(
        client, db_session, org_id, token, "q.xlsx", _build_xlsx_bytes(), "Yes."
    )

    resp = await client.get(
        f"/api/orgs/{org_id}/projects/{project_id}/export?format=docx", headers=_auth(token)
    )
    assert resp.status_code == 400


async def test_export_tenant_isolation(client, db_session):
    token_a = await _signup(client, "orga@example.com")
    org_a = await _create_org(client, token_a, "Org A")
    project_id = await _setup_project_with_answer(
        client, db_session, org_a, token_a, "q.xlsx", _build_xlsx_bytes(), "Yes."
    )

    token_b = await _signup(client, "orgb@example.com")
    org_b = await _create_org(client, token_b, "Org B")

    resp = await client.get(
        f"/api/orgs/{org_a}/projects/{project_id}/export?format=xlsx", headers=_auth(token_b)
    )
    assert resp.status_code == 403

    resp = await client.get(
        f"/api/orgs/{org_b}/projects/{project_id}/export?format=xlsx", headers=_auth(token_b)
    )
    assert resp.status_code == 404
